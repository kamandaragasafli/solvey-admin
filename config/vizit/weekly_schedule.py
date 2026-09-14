from datetime import date, timedelta
from io import BytesIO
from pathlib import Path
import re

import openpyxl
from django.contrib import messages
from django.core import signing
from django.http import Http404, HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.utils import timezone
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib.colors import HexColor, white
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas

from regions.models import Region

from .models import (
    Istifadeci,
    WeeklySchedule,
    WeeklyScheduleDay,
    WeeklyScheduleVisit,
)
from .utils import vizit_login_required

_PDF_FONTS_READY = False
_SHARE_SALT = "vizit-weekly-schedule-pdf"
_SHARE_MAX_AGE = 60 * 60 * 24 * 30  # 30 gün


def _schedule_pdf_filename(schedule):
    """PDF adı: «Username Həftəlik Qrafiq.pdf»."""
    name = ""
    user = getattr(schedule, "created_by", None)
    if user is not None:
        name = (user.ad or user.login or "").strip()
    if not name:
        name = "Qrafiq"
    safe = re.sub(r'[\\/:*?"<>|]+', "", name).strip()
    safe = re.sub(r"\s+", " ", safe) or "Qrafiq"
    return f"{safe} Həftəlik Qrafiq.pdf"


def _pdf_content_disposition(filename, disposition="inline"):
    from urllib.parse import quote

    ascii_name = filename.encode("ascii", "ignore").decode("ascii") or "qrafiq.pdf"
    utf8_name = quote(filename)
    return (
        f"{disposition}; filename=\"{ascii_name}\"; "
        f"filename*=UTF-8''{utf8_name}"
    )


def _ensure_pdf_fonts():
    """Windows Arial — Azərbaycan hərfləri üçün."""
    global _PDF_FONTS_READY
    if _PDF_FONTS_READY:
        return
    fonts = Path(r"C:\Windows\Fonts")
    regular = fonts / "arial.ttf"
    bold = fonts / "arialbd.ttf"
    if not regular.exists():
        # Linux fallback
        regular = Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf")
        bold = Path("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf")
    pdfmetrics.registerFont(TTFont("WQ", str(regular)))
    pdfmetrics.registerFont(TTFont("WQ-Bold", str(bold if bold.exists() else regular)))
    _PDF_FONTS_READY = True


def _bolgeler_for_user(user_rol, user_bolge_ids):
    if user_rol in (Istifadeci.ROL_REHBER, Istifadeci.ROL_DIVIZIYA_REHB):
        return Region.objects.order_by("region_name")
    return Region.objects.filter(pk__in=user_bolge_ids).order_by("region_name")


WEEKDAY_LABELS = [
    (1, "Bazar ertəsi"),
    (2, "Çərşənbə axşamı"),
    (3, "Çərşənbə"),
    (4, "Cümə axşamı"),
    (5, "Cümə"),
]


def _monday_of(d):
    return d - timedelta(days=d.weekday())


def _schedule_days_payload(schedule):
    by_wd = {day.weekday: day for day in schedule.days.all()}
    rows = []
    for wd, label in WEEKDAY_LABELS:
        day = by_wd.get(wd)
        visits = list(day.visits.all()) if day else []
        rows.append(
            {
                "weekday": wd,
                "label": label,
                "day": day,
                "region_id": day.region_id if day else "",
                "region_name": day.region.region_name if day and day.region_id else "",
                "qrafik": day.qrafik if day else "",
                "tn": day.tn if day else "",
                "visits": visits,
                "visit_lines": "\n".join(v.place_name for v in visits),
            }
        )
    return rows


def _save_weekly_schedule_from_post(request, schedule):
    schedule.menecer_name = (request.POST.get("menecer_name") or "").strip()
    schedule.sedr_name = (request.POST.get("sedr_name") or "").strip()
    schedule.note = (request.POST.get("note") or "").strip()
    schedule.save()

    for wd, _label in WEEKDAY_LABELS:
        region_raw = (request.POST.get(f"region_{wd}") or "").strip()
        qrafik = (request.POST.get(f"qrafik_{wd}") or "").strip()
        tn = (request.POST.get(f"tn_{wd}") or "").strip()
        visits_raw = request.POST.get(f"visits_{wd}") or ""
        region_id = int(region_raw) if region_raw.isdigit() else None

        day, _ = WeeklyScheduleDay.objects.update_or_create(
            schedule=schedule,
            weekday=wd,
            defaults={
                "region_id": region_id,
                "qrafik": qrafik,
                "tn": tn,
            },
        )
        day.visits.all().delete()
        lines = [ln.strip() for ln in visits_raw.replace("\r", "").split("\n") if ln.strip()]
        WeeklyScheduleVisit.objects.bulk_create(
            [
                WeeklyScheduleVisit(day=day, place_name=name, position=i)
                for i, name in enumerate(lines)
            ]
        )


@vizit_login_required
def weekly_schedule_list(request):
    user_id = request.session.get("istifadeci_id")
    user_rol = request.session.get("rol")
    qs = WeeklySchedule.objects.select_related("created_by").prefetch_related("days__region")
    if user_rol == Istifadeci.ROL_NUMAYENDE:
        qs = qs.filter(created_by_id=user_id)
    return render(request, "vizit/weekly-schedule-list.html", {"schedules": qs[:50]})


@vizit_login_required
def weekly_schedule_create(request):
    user_rol = request.session.get("rol")
    user_bolge_ids = request.session.get("bolge_ids", [])
    regions = _bolgeler_for_user(user_rol, user_bolge_ids)
    default_monday = _monday_of(timezone.localdate())

    if request.method == "POST":
        week_raw = (request.POST.get("week_start") or "").strip()
        try:
            week_start = _monday_of(date.fromisoformat(week_raw))
        except ValueError:
            week_start = default_monday
        schedule = WeeklySchedule.objects.create(
            week_start=week_start,
            created_by_id=request.session.get("istifadeci_id"),
        )
        _save_weekly_schedule_from_post(request, schedule)
        messages.success(request, "Həftəlik qrafiq yadda saxlanıldı.")
        return redirect("vizit:weekly_schedule_detail", pk=schedule.pk)

    empty_days = [
        {
            "weekday": wd,
            "label": label,
            "region_id": "",
            "qrafik": "",
            "tn": "",
            "visit_lines": "",
        }
        for wd, label in WEEKDAY_LABELS
    ]
    return render(
        request,
        "vizit/weekly-schedule-form.html",
        {
            "regions": regions,
            "days": empty_days,
            "week_start": default_monday.isoformat(),
            "menecer_name": "",
            "sedr_name": "",
            "note": "",
            "is_edit": False,
            "schedule": None,
        },
    )


@vizit_login_required
def weekly_schedule_edit(request, pk):
    schedule = get_object_or_404(
        WeeklySchedule.objects.prefetch_related("days__visits", "days__region"),
        pk=pk,
    )
    user_id = request.session.get("istifadeci_id")
    user_rol = request.session.get("rol")
    if user_rol == Istifadeci.ROL_NUMAYENDE and schedule.created_by_id != user_id:
        messages.error(request, "Bu qrafiki redaktə etmək icazəniz yoxdur.")
        return redirect("vizit:weekly_schedule_list")

    regions = _bolgeler_for_user(user_rol, request.session.get("bolge_ids", []))
    if request.method == "POST":
        week_raw = (request.POST.get("week_start") or "").strip()
        try:
            schedule.week_start = _monday_of(date.fromisoformat(week_raw))
        except ValueError:
            pass
        _save_weekly_schedule_from_post(request, schedule)
        messages.success(request, "Həftəlik qrafiq yeniləndi.")
        return redirect("vizit:weekly_schedule_detail", pk=schedule.pk)

    return render(
        request,
        "vizit/weekly-schedule-form.html",
        {
            "regions": regions,
            "days": _schedule_days_payload(schedule),
            "week_start": schedule.week_start.isoformat(),
            "menecer_name": schedule.menecer_name,
            "sedr_name": schedule.sedr_name,
            "note": schedule.note,
            "is_edit": True,
            "schedule": schedule,
        },
    )


@vizit_login_required
def weekly_schedule_detail(request, pk):
    schedule = get_object_or_404(
        WeeklySchedule.objects.select_related("created_by").prefetch_related(
            "days__visits", "days__region"
        ),
        pk=pk,
    )
    token = signing.dumps({"id": schedule.pk}, salt=_SHARE_SALT)
    share_path = reverse("vizit:weekly_schedule_shared_pdf", args=[token])
    share_url = request.build_absolute_uri(share_path)
    return render(
        request,
        "vizit/weekly-schedule-detail.html",
        {
            "schedule": schedule,
            "days": _schedule_days_payload(schedule),
            "share_url": share_url,
            "pdf_filename": _schedule_pdf_filename(schedule),
        },
    )


@vizit_login_required
def weekly_schedule_delete(request, pk):
    schedule = get_object_or_404(WeeklySchedule, pk=pk)
    user_id = request.session.get("istifadeci_id")
    user_rol = request.session.get("rol")
    if user_rol == Istifadeci.ROL_NUMAYENDE and schedule.created_by_id != user_id:
        messages.error(request, "Bu qrafiki silmək icazəniz yoxdur.")
        return redirect("vizit:weekly_schedule_list")
    schedule.delete()
    messages.success(request, "Həftəlik qrafiq silindi.")
    return redirect("vizit:weekly_schedule_list")


@vizit_login_required
def weekly_schedule_excel(request, pk):
    schedule = get_object_or_404(
        WeeklySchedule.objects.prefetch_related("days__visits", "days__region"),
        pk=pk,
    )
    days = _schedule_days_payload(schedule)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Həftəlik Qrafiq"

    # Kağız blanka yaxın rənglər
    navy = "1A5276"
    soft = "D6EAF8"
    soft2 = "EBF5FB"
    grid = "1A5276"
    ink = "1C2833"

    header_fill = PatternFill(start_color=navy, end_color=navy, fill_type="solid")
    meta_fill = PatternFill(start_color=soft, end_color=soft, fill_type="solid")
    alt_fill = PatternFill(start_color=soft2, end_color=soft2, fill_type="solid")
    white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

    title_font = Font(name="Calibri", bold=True, color="FFFFFF", size=14)
    day_font = Font(name="Calibri", bold=True, color="FFFFFF", size=12)
    meta_font = Font(name="Calibri", bold=True, color=ink, size=11)
    body_font = Font(name="Calibri", color=ink, size=11)
    sign_font = Font(name="Calibri", bold=True, color=ink, size=11)

    thin = Border(
        left=Side(style="thin", color=grid),
        right=Side(style="thin", color=grid),
        top=Side(style="thin", color=grid),
        bottom=Side(style="thin", color=grid),
    )
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left_mid = Alignment(horizontal="left", vertical="center", wrap_text=True)

    # —— Başlıq ——
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=5)
    for col in range(1, 6):
        c = ws.cell(row=1, column=col)
        c.fill = header_fill
        c.border = thin
    title = ws.cell(
        row=1,
        column=1,
        value=f"Həftəlik qrafiq — {schedule.week_start:%d.%m.%Y} – {schedule.week_end:%d.%m.%Y}",
    )
    title.font = title_font
    title.alignment = center
    ws.row_dimensions[1].height = 28

    # —— Gün başlıqları + Bölgə + Qrafik ——
    for col, day in enumerate(days, 1):
        head = ws.cell(row=2, column=col, value=day["label"])
        head.fill = header_fill
        head.font = day_font
        head.alignment = center
        head.border = thin

        bolge = ws.cell(
            row=3,
            column=col,
            value=f"Bölgə: {day['region_name'] or '—'}",
        )
        bolge.fill = meta_fill
        bolge.font = meta_font
        bolge.alignment = left_mid
        bolge.border = thin

        qraf = ws.cell(
            row=4,
            column=col,
            value=f"Qrafik: {day['qrafik'] or '—'}",
        )
        qraf.fill = meta_fill
        qraf.font = meta_font
        qraf.alignment = left_mid
        qraf.border = thin

    ws.row_dimensions[2].height = 24
    ws.row_dimensions[3].height = 22
    ws.row_dimensions[4].height = 22

    # —— Vizit yerləri (minimum 8 sətir — blank kimi) ——
    max_visits = max(max((len(d["visits"]) for d in days), default=0), 8)
    for i in range(max_visits):
        row = 5 + i
        ws.row_dimensions[row].height = 20
        for col, day in enumerate(days, 1):
            name = day["visits"][i].place_name if i < len(day["visits"]) else ""
            cell = ws.cell(row=row, column=col, value=name)
            cell.fill = alt_fill if i % 2 else white_fill
            cell.font = body_font
            cell.border = thin
            cell.alignment = left_mid

    # —— T/N ——
    tn_row = 5 + max_visits
    ws.row_dimensions[tn_row].height = 22
    for col, day in enumerate(days, 1):
        cell = ws.cell(row=tn_row, column=col, value=f"T/N: {day['tn'] or ''}")
        cell.fill = meta_fill
        cell.font = meta_font
        cell.border = thin
        cell.alignment = left_mid

    # —— İmza sətirləri ——
    sign_row = tn_row + 2
    ws.row_dimensions[sign_row].height = 24
    ws.merge_cells(start_row=sign_row, start_column=1, end_row=sign_row, end_column=2)
    menecer = ws.cell(
        row=sign_row,
        column=1,
        value=f"Menecer: {schedule.menecer_name or ''}",
    )
    menecer.font = sign_font
    menecer.alignment = left_mid

    ws.merge_cells(start_row=sign_row, start_column=3, end_row=sign_row, end_column=5)
    sedr = ws.cell(
        row=sign_row,
        column=3,
        value=f"İdarə heyətinin sədri: {schedule.sedr_name or ''}",
    )
    sedr.font = sign_font
    sedr.alignment = left_mid

    for col in range(1, 6):
        ws.column_dimensions[get_column_letter(col)].width = 26

    # Çap: A4 landşaft, 1 səhifəyə sığsın
    ws.print_area = f"A1:E{sign_row}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.page_margins.left = 0.4
    ws.page_margins.right = 0.4
    ws.page_margins.top = 0.5
    ws.page_margins.bottom = 0.5
    ws.sheet_view.showGridLines = False

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    filename = f"heftelik_qrafiq_{schedule.week_start.isoformat()}.xlsx"
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    wb.save(response)
    return response


def _draw_wrapped_text(c, text, x, y, max_width, font="WQ", size=8, leading=11, color=None):
    """Sadə söz-wrap; istifadə olunan hündürlüyü qaytarır."""
    c.setFont(font, size)
    if color:
        c.setFillColor(color)
    words = (text or "").split()
    if not words:
        return 0
    lines = []
    current = words[0]
    for w in words[1:]:
        trial = f"{current} {w}"
        if c.stringWidth(trial, font, size) <= max_width:
            current = trial
        else:
            lines.append(current)
            current = w
    lines.append(current)
    for i, line in enumerate(lines):
        c.drawString(x, y - i * leading, line)
    return len(lines) * leading


@vizit_login_required
def weekly_schedule_pdf(request, pk):
    """Ekrandakı 5 sütunlu kart layout-unun eyni PDF-i (A4 landşaft)."""
    schedule = get_object_or_404(
        WeeklySchedule.objects.select_related("created_by").prefetch_related(
            "days__visits", "days__region"
        ),
        pk=pk,
    )
    pdf_bytes, filename = _build_weekly_schedule_pdf(schedule)
    response = HttpResponse(pdf_bytes, content_type="application/pdf")
    as_attachment = request.GET.get("download") == "1"
    disp = "attachment" if as_attachment else "inline"
    response["Content-Disposition"] = _pdf_content_disposition(filename, disp)
    return response


def weekly_schedule_shared_pdf(request, token):
    """Login olmadan açılan müvəqqəti paylaşım linki (30 gün)."""
    try:
        data = signing.loads(token, salt=_SHARE_SALT, max_age=_SHARE_MAX_AGE)
        pk = int(data["id"])
    except (signing.BadSignature, signing.SignatureExpired, KeyError, TypeError, ValueError):
        raise Http404("Paylaşım linki etibarsızdır və ya müddəti bitib.")
    schedule = get_object_or_404(
        WeeklySchedule.objects.select_related("created_by").prefetch_related(
            "days__visits", "days__region"
        ),
        pk=pk,
    )
    pdf_bytes, filename = _build_weekly_schedule_pdf(schedule)
    response = HttpResponse(pdf_bytes, content_type="application/pdf")
    response["Content-Disposition"] = _pdf_content_disposition(filename, "inline")
    return response


def _build_weekly_schedule_pdf(schedule):
    days = _schedule_days_payload(schedule)
    _ensure_pdf_fonts()

    navy = HexColor("#1A5276")
    meta_bg = HexColor("#F0F7FF")
    tn_bg = HexColor("#F8FAFC")
    line = HexColor("#DCE6F0")
    dashed = HexColor("#E8EDF2")
    ink = HexColor("#1C2833")
    muted = HexColor("#94A3B8")

    buffer = BytesIO()
    page_w, page_h = landscape(A4)
    c = canvas.Canvas(buffer, pagesize=landscape(A4))

    margin_x = 10 * mm
    margin_y = 10 * mm
    gap = 3 * mm
    cols = 5
    usable_w = page_w - 2 * margin_x
    col_w = (usable_w - gap * (cols - 1)) / cols

    title = (
        f"Həftəlik qrafiq — {schedule.week_start:%d.%m.%Y} – {schedule.week_end:%d.%m.%Y}"
    )
    c.setFillColor(navy)
    c.setFont("WQ-Bold", 13)
    c.drawCentredString(page_w / 2, page_h - margin_y - 4 * mm, title)

    top = page_h - margin_y - 12 * mm
    bottom = margin_y + 14 * mm
    col_h = top - bottom

    head_h = 9 * mm
    meta_h = 7 * mm
    tn_h = 8 * mm

    for i, day in enumerate(days):
        x = margin_x + i * (col_w + gap)
        y = bottom

        c.setStrokeColor(navy)
        c.setLineWidth(1)
        c.setFillColor(white)
        c.roundRect(x, y, col_w, col_h, 4, stroke=1, fill=1)

        c.setFillColor(navy)
        c.rect(x, y + col_h - head_h, col_w, head_h, stroke=0, fill=1)
        c.setFont("WQ-Bold", 9)
        c.setFillColor(white)
        c.drawCentredString(x + col_w / 2, y + col_h - head_h + 3.2 * mm, day["label"])

        meta_top = y + col_h - head_h
        for idx, (label, value) in enumerate(
            [
                ("Bölgə:", day["region_name"] or "—"),
                ("Qrafik:", day["qrafik"] or "—"),
            ]
        ):
            my = meta_top - (idx + 1) * meta_h
            c.setFillColor(meta_bg)
            c.rect(x + 0.4, my, col_w - 0.8, meta_h, stroke=0, fill=1)
            c.setStrokeColor(line)
            c.setLineWidth(0.4)
            c.line(x, my, x + col_w, my)
            c.setFillColor(ink)
            c.setFont("WQ-Bold", 8)
            label_w = c.stringWidth(label + " ", "WQ-Bold", 8)
            c.drawString(x + 2 * mm, my + 2.2 * mm, label)
            c.setFont("WQ", 8)
            max_val_w = col_w - 4 * mm - label_w
            val = value
            while c.stringWidth(val, "WQ", 8) > max_val_w and len(val) > 3:
                val = val[:-2] + "…"
            c.drawString(x + 2 * mm + label_w, my + 2.2 * mm, val)

        visits_top = meta_top - 2 * meta_h
        visits_bottom = y + tn_h
        pad = 2.2 * mm
        text_x = x + pad
        text_y = visits_top - 4 * mm
        max_text_w = col_w - 2 * pad
        visits = day["visits"]
        if not visits:
            c.setFillColor(muted)
            c.setFont("WQ", 8)
            c.drawString(text_x, text_y, "—")
        else:
            for vi, v in enumerate(visits, 1):
                if text_y < visits_bottom + 3 * mm:
                    break
                line_txt = f"{vi}. {v.place_name}"
                used = _draw_wrapped_text(
                    c, line_txt, text_x, text_y, max_text_w, font="WQ", size=8, leading=10, color=ink
                )
                sep_y = text_y - max(used, 10) + 2
                if sep_y > visits_bottom + 2 * mm:
                    c.setStrokeColor(dashed)
                    c.setDash(1, 2)
                    c.setLineWidth(0.5)
                    c.line(x + pad, sep_y, x + col_w - pad, sep_y)
                    c.setDash()
                text_y -= max(used, 10) + 2

        c.setFillColor(tn_bg)
        c.rect(x + 0.4, y + 0.4, col_w - 0.8, tn_h - 0.4, stroke=0, fill=1)
        c.setStrokeColor(line)
        c.setLineWidth(0.5)
        c.line(x, y + tn_h, x + col_w, y + tn_h)
        c.setFillColor(ink)
        c.setFont("WQ-Bold", 8)
        tn_label = "T/N:"
        c.drawString(x + 2 * mm, y + 2.8 * mm, tn_label)
        c.setFont("WQ", 8)
        c.drawString(
            x + 2 * mm + c.stringWidth(tn_label + " ", "WQ-Bold", 8),
            y + 2.8 * mm,
            day["tn"] or "",
        )

        c.setStrokeColor(navy)
        c.setLineWidth(1.2)
        c.roundRect(x, y, col_w, col_h, 4, stroke=1, fill=0)

    c.setFillColor(ink)
    c.setFont("WQ-Bold", 10)
    c.drawString(
        margin_x,
        margin_y + 3 * mm,
        f"Menecer: {schedule.menecer_name or ''}",
    )
    sedr = f"İdarə heyətinin sədri: {schedule.sedr_name or ''}"
    c.drawRightString(page_w - margin_x, margin_y + 3 * mm, sedr)

    c.showPage()
    c.save()
    buffer.seek(0)
    return buffer.getvalue(), _schedule_pdf_filename(schedule)

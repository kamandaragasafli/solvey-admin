"""
1C-dən ixrac olunan "Продажи" pivot hesabatını (Номенклатура -> Bölgə -> Şəhər -> Aptek)
oxuyub, hər dərman üçün {şəhər: say} formasına salır.

Sətir iyerarxiyası Excel-in "outline level" (qruplaşdırma) məlumatı ilə oxunur:
    outline 0  -> dərman (Номенклатура) sətri
    outline 1  -> region sətri: ya "Bakı" kimi son-nöqtə şəhər, ya da "Bölgə (X Depo)"
    outline 2  -> "Bölgə (...)" daxilindəki əsl şəhər / rayon
    outline 3  -> aptek (cəmə artıq daxildir)
"""
from __future__ import annotations

import openpyxl

from depo_sales.constants import (
    BAKU_SOURCE_KEYS,
    CITY_NAME_ALIASES,
    MERGED_BAKU_COLUMN,
)

DATA_START_ROW = 14
DRUG_COLUMN = 2
VALUE_COLUMN = 3
TOTAL_ROW_LABEL = "Итог"


def parse_sales_workbook(file_obj_or_path) -> dict[str, dict[str, int]]:
    """
    Excel faylını (path və ya file-like) oxuyub {dərman_adı: {şəhər: say}} qaytarır.
    Bakı və Abşeron avtomatik olaraq tək "Bakı+Abşeron" sütununda birləşdirilir.
    """
    wb = openpyxl.load_workbook(file_obj_or_path, data_only=True)
    ws = wb.active

    drugs: dict[str, dict[str, int]] = {}
    current_drug = None
    current_region = None

    for row_idx in range(DATA_START_ROW, ws.max_row + 1):
        label = ws.cell(row=row_idx, column=DRUG_COLUMN).value
        value = ws.cell(row=row_idx, column=VALUE_COLUMN).value
        outline = ws.row_dimensions[row_idx].outlineLevel

        if label is None:
            continue

        if outline == 0:
            if label == TOTAL_ROW_LABEL:
                current_drug = None
                continue
            current_drug = label.strip()
            drugs.setdefault(current_drug, {})

        elif outline == 1 and current_drug:
            current_region = label
            if not label.startswith("Bölgə"):
                qty = value or 0
                drugs[current_drug][label] = drugs[current_drug].get(label, 0) + qty

        elif outline == 2 and current_drug:
            if current_region and current_region.startswith("Bölgə"):
                qty = value or 0
                drugs[current_drug][label] = drugs[current_drug].get(label, 0) + qty

    _normalize_city_names(drugs)
    _merge_baku_and_absheron(drugs)
    return drugs


def _normalize_city_names(drugs: dict[str, dict[str, int]]) -> None:
    for city_qty in drugs.values():
        for wrong, correct in CITY_NAME_ALIASES.items():
            if wrong in city_qty:
                city_qty[correct] = city_qty.get(correct, 0) + city_qty.pop(wrong)


def _merge_baku_and_absheron(drugs: dict[str, dict[str, int]]) -> None:
    for city_qty in drugs.values():
        total = 0
        found = False
        for key in BAKU_SOURCE_KEYS:
            if key in city_qty:
                total += city_qty.pop(key)
                found = True
        if found:
            city_qty[MERGED_BAKU_COLUMN] = city_qty.get(MERGED_BAKU_COLUMN, 0) + total
